# -*- coding: utf-8 -*-
"""
投标文件「人员简历」独立拆分与文本抽取（复用 doc_splitter/splitter.py 的标题识别与 docx 切割逻辑）

流程简述：
  1. 扫描 outlineLevel=0 的一级标题，按关键词筛出「简历/团队/人员」等相关章节
  2. 调用 split_docx 仅导出这些一级章节为独立 docx
  3. 对每个一级 docx 再按 outlineLevel=1 切二级块（常见为单人简历）
  4. 对每个文本块做轻量正则抽取（姓名/职务/学历等），写入 manifest.json
    5. 对含「角色-姓名」或「编号小节 + 姓名」（如 6.2.1. 张云聪）换行结构的正文再拆 person_entries

用法：
  python resume_split_extract.py <输入.docx> [路径相关选项，见 --help]

输出子目录默认：{文件夹主名}__{标签}_{批次时间}/
  主名由 --folder-base-from / --folder-base 决定；标签默认「简历解析」可用 --folder-tag 改。
  环境变量（未传 CLI 时作默认）：RESUME_EXTRACT_OUT、RESUME_EXTRACT_FOLDER_BASE_FROM、
  RESUME_EXTRACT_FOLDER_BASE、RESUME_EXTRACT_FOLDER_TAG、RESUME_EXTRACT_BATCH_ID。

依赖：与 doc_splitter 相同（python-docx, lxml）
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import sys
from datetime import datetime, timezone
import unicodedata
from pathlib import Path

_THIS = Path(__file__).resolve().parent
_SERVICES = _THIS.parent
_SPLITTER_PATH = _SERVICES / "doc_splitter" / "splitter.py"


def _load_splitter():
    spec = importlib.util.spec_from_file_location("bid_doc_splitter", _SPLITTER_PATH)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(mod)
    return mod


_splitter = _load_splitter()
split_docx = _splitter.split_docx


def _load_sibling_module(mod_name: str, file_name: str):
    p = _THIS / file_name
    spec = importlib.util.spec_from_file_location(mod_name, p)
    m = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(m)
    return m


_enrich = _load_sibling_module("resume_ai_enrich", "resume_ai_enrich.py")
enrich_person_row = _enrich.enrich_person_row
enrich_section_extracted = _enrich.enrich_section_extracted

# 一级章节标题命中任一关键词即视为「可能含人员简历」的章节（可按项目再调）
RESUME_H1_KEYWORDS = (
    "简历",
    "人员",
    "团队",
    "拟投入",
    "项目成员",
    "项目班子",
    "技术人员",
    "主要人员",
    "人员表",
    "人员情况",
    "人员配置",
    "团队人员",
    "人员资格",
    "拟派",
    "项目负责人",
    "技术负责人",
)

# 二级标题补充（避免「公司人员结构」等纯概述小节误命中）
RESUME_H2_EXTRA_KEYWORDS = (
    "项目组织",
    "项目团队",
    "人员一览",
    "人员汇总",
    "人员情况表",
    "拟任人员",
    "人员配备",
    "人员资质",
    "项目组成员",
    "开发团队",
    "服务团队",
    "人员简历",
    "人员清单",
    "拟派人员",
)

# 仅标题为下列之一时不作为简历模块（常为概述、非逐人简历）
RESUME_H2_BLOCKLIST = frozenset(
    {
        "公司人员结构",
        "公司基本情况",
        "资格证明文件",
    }
)


def normalize_title(t: str) -> str:
    return (t or "").replace(" ", "").replace("\u3000", "").strip()


def h1_matches_resume(title: str) -> bool:
    n = normalize_title(title)
    if not n:
        return False
    return any(k in n for k in RESUME_H1_KEYWORDS)


def _expand_h2_sections_with_next_heading(
    all_h2: list[tuple[int, str]], resume_sections: list[dict]
) -> tuple[list[dict], set[str]]:
    """
    为 split_docx 构造章节列表：在每个待导出的 H2 后追加文档中的「下一个 H2」标题作切割边界，
    避免仅一节时 end_idx 落到全文末尾。返回 (expanded_sections, 边界标题集合) 边界标题对应文件不入库。
    """
    skip: set[str] = set()
    if not resume_sections or not all_h2:
        return resume_sections, skip

    titles_in_order = [s["title"] for s in resume_sections]
    # 建立标题 → 在 all_h2 中首次出现的位置
    pos_by_title: dict[str, int] = {}
    for j, (_ix, t) in enumerate(all_h2):
        if t not in pos_by_title:
            pos_by_title[t] = j

    expanded: list[dict] = []
    for wi, want in enumerate(titles_in_order):
        expanded.append({"title": want, "start_keyword": want})
        pos = pos_by_title.get(want)
        if pos is None:
            continue
        if pos + 1 >= len(all_h2):
            continue
        nxt_title = all_h2[pos + 1][1]
        next_wanted = titles_in_order[wi + 1] if wi + 1 < len(titles_in_order) else None
        if next_wanted is not None and nxt_title == next_wanted:
            continue
        expanded.append({"title": nxt_title, "start_keyword": nxt_title})
        skip.add(nxt_title)

    return expanded, skip


def h2_matches_resume(title: str) -> bool:
    n = normalize_title(title)
    if not n:
        return False
    if n in RESUME_H2_BLOCKLIST:
        return False
    keys = RESUME_H1_KEYWORDS + RESUME_H2_EXTRA_KEYWORDS
    if any(k in n for k in keys):
        return True
    # 「组织架构」需与项目/简历语境同时出现，避免泛泛章节
    if "组织架构" in n and ("项目" in n or "简历" in n or "人员" in n):
        return True
    return False


def extract_fields_from_text(text: str, fallback_title: str = "") -> dict:
    """轻量规则抽取（不调用大模型）；后续可接 AI 覆盖 extracted_json。"""
    t = _normalize_resume_text(text or "")
    out = {
        "person_name": None,
        "role_title": None,
        "education_hint": None,
        "cert_hint": None,
        "raw_preview": t[:2000],
    }

    def _m(pattern: str, flags=0):
        m = re.search(pattern, t, flags)
        return m.group(1).strip() if m else None

    name = _m(r"姓\s*名\s*[:：]\s*([^\s\n，,；;]{1,20})")
    if not name:
        name = _m(r"姓名\s*[:：]\s*([^\s\n，,；;]{1,20})")
    name = _sanitize_person_name_string(name)
    if not name:
        name = _extract_name_near_xingming(t)
    out["person_name"] = name

    role = _m(r"(?:拟任职务|职务|岗位|职位)\s*[:：]\s*([^\n]{1,80})")
    out["role_title"] = role

    edu = re.search(
        r"(本科|硕士|博士|专科|学士|研究生).{0,40}(?:毕业|学历|院校|专业)",
        t,
    )
    if edu:
        s = max(0, edu.start() - 20)
        e = min(len(t), edu.end() + 40)
        out["education_hint"] = t[s:e].replace("\n", " ").strip()

    cert = re.search(r"(?:证书|认证|PMP|软考|职称).{0,80}", t)
    if cert:
        out["cert_hint"] = cert.group(0).replace("\n", " ").strip()

    if not out["person_name"] and fallback_title:
        ft = fallback_title.strip()
        sub_role, sub_name = _parse_title_colon_trailing_name(ft)
        if sub_name:
            out["person_name"] = sub_name
            if sub_role and not out["role_title"]:
                out["role_title"] = sub_role
        elif re.match(r"^（\d+）", ft) or re.match(r"^\d+\.\d+", ft):
            pass
        elif (
            1 < len(ft) <= 20
            and not any(x in ft for x in ("章", "节", "部分", "附件", "证书", "证明"))
            and _is_plausible_cn_person_name(ft)
        ):
            out["person_name"] = ft

    return out


# 绝不能当作自然人姓名的字段标签（表格串行、合并单元格导出时常夹在「姓名」与「性别」之间）
_INVALID_EXACT_PERSON_NAME = frozenset(
    {
        "身份证",
        "身份证号",
        "性别",
        "姓名",
        "民族",
        "籍贯",
        "学历",
        "专业",
        "年龄",
        "电话",
        "手机",
        "邮箱",
        "职务",
        "岗位",
        "职位",
        "人员级别",
        "出生年月",
        "出生日期",
        "毕业院校",
        "毕业学校",
        "公司岗位",
        "拟在本项目担任中职务",
        "职称",
    }
)


def _looks_like_cn_id_card(s: str) -> bool:
    x = re.sub(r"\s+", "", (s or "").strip())
    if not x:
        return False
    return bool(re.fullmatch(r"[0-9Xx\*]{15,22}", x))


def _sanitize_person_name_string(name: str | None) -> str | None:
    """姓名候选：去掉标签词、身份证号误当姓名。"""
    if not name:
        return None
    n = (name or "").strip()
    if not n:
        return None
    n_plain = re.sub(r"[\s\u3000]+", "", n)
    if n_plain in _INVALID_EXACT_PERSON_NAME:
        return None
    if _looks_like_cn_id_card(n_plain):
        return None
    if len(n_plain) >= 15 and re.fullmatch(r"[0-9Xx\*]+", n_plain):
        return None
    return n


def _extract_name_near_xingming(t: str) -> str | None:
    """
    表格导出错位时，优先用「姓名…性别」夹逼的真实姓名；
    其次「姓名\\t真实名\\t」制表行。
    """
    if not t:
        return None
    m = re.search(
        r"姓名\s*[：:\t]\s*([\u4e00-\u9fa5·]{2,10})\s*(?:[：:\t]?\s*性别|[\t\n\r]|(?=[\u4e00-\u9fa5]{2,4}\s*[：:\t]))",
        t,
    )
    if m:
        cand = m.group(1).strip()
        if _is_plausible_cn_person_name(cand):
            return cand[:40]
    m2 = re.search(r"姓名\s*[：:\t]\s*([\u4e00-\u9fa5·]{2,10})\t", t)
    if m2:
        cand = m2.group(1).strip()
        if _is_plausible_cn_person_name(cand):
            return cand[:40]
    return None


def _parse_title_colon_trailing_name(title: str) -> tuple[str | None, str | None]:
    """「（一）项目总监：贾永超」「基础人员：张雅茹」→ (职务片段, 姓名)。"""
    s = (title or "").strip()
    if not s:
        return None, None
    m = re.search(
        r"(?:[（(][一二三四五六七八九十百千]+[）)]\s*)?([^：:\n]{1,40})\s*[：:]\s*([\u4e00-\u9fa5·]{2,10})\s*$",
        s,
    )
    if not m:
        return None, None
    role = re.sub(r"^[（(][一二三四五六七八九十百千]+[）)]\s*", "", m.group(1).strip())
    name = m.group(2).strip()
    if _is_plausible_cn_person_name(name):
        return (role[:120] if role else None, name[:40])
    return None, None


# 常见「角色-姓名」独立成行，或「3.2.1. 项目经理-崔学佳」编号+职务+姓名（职务段放宽防截断）
_PERSON_HEADER = re.compile(
    r"(?:^|\n)"
    r"(?P<role>[\u4e00-\u9fa5A-Za-z0-9\、\.\+\s（）()]{2,80})"
    r"[-－]"
    r"(?P<name>[\u4e00-\u9fa5·]{2,10})"
    r"\s*\n"
)

# 制表/换行混排时，小节标题可能紧跟在 \t 后且与上一段同一物理行（无换行）
_PERSON_HEADER_INLINE = re.compile(
    r"(?:^|[\n\r\t])"
    r"(?P<role>\d{1,2}(?:\.\d{1,3})+\.?\s+[\u4e00-\u9fa5A-Za-z0-9（）()]{2,40})"
    r"[-－]"
    r"(?P<name>[\u4e00-\u9fa5·]{2,10})"
    r"(?=[\n\r\t\u3000 ]|\Z)"
)

# 投标文件中常见「6.2.1. 张云聪」小节标题（非「职务-姓名」）；须配合块内简历表头再落人，防误切技术小节
_PERSON_HEADER_NUMBERED = re.compile(
    r"(?:^|\n)"
    r"\s*"
    r"(?P<prefix>\d+(?:\.\d+)+)"
    r"\.?"
    r"\s+"
    r"(?P<name>[\u4e00-\u9fa5·]{2,10})"
    r"\s*"
    r"(?:\n|\r|\t|\Z)"
)

# 二级切片 docx 导出后常见首行仅姓名，下一行起为「姓名\\t…」表格转文本（与 Word 里 6.2.1 标题分离）
_PERSON_HEADER_LEADING_NAME = re.compile(
    r"^(?P<name>[\u4e00-\u9fa5·]{2,10})\s*\n"
    r"(?=[\s\S]{0,600}(?:姓名\s*[：:\t]|性别\s*[：:\t]))"
)

# 「（一）项目总监：贾永超」「基础人员：张雅茹」——中文投标简历常见小节标题（非连字符）
_PERSON_HEADER_ENUM_COLON = re.compile(
    r"(?:^|\n)"
    r"\s*(?:[（(][一二三四五六七八九十百千]+[）)]\s*)?"
    r"(?P<role>[^\n\r：:]{2,48})"
    r"[：:]\s*"
    r"(?P<name>[\u4e00-\u9fa5·]{2,10})",
)

# 单行「角色-姓名」（与 _PERSON_HEADER 一致，用于按段落实体顺序累计当前人）
_PERSON_LINE_HEADER = re.compile(
    r"^([\u4e00-\u9fa5A-Za-z0-9\、\.\+\s（）()]{2,44})[-－]([\u4e00-\u9fa5·]{2,10})\s*$"
)

# 误当成「姓名」的常见需求/设计书用语（投标技术册与人员简历混排时）
_RESUME_NAME_DENY_SUBSTR = (
    "组件",
    "规则",
    "校验",
    "修改",
    "内容",
    "需求",
    "响应",
    "数据",
    "账户",
    "信息",
    "报表",
    "自动化",
    "取数",
    "功能",
    "模块",
    "接口",
    "字段",
    "系统",
    "存款",
    "合约",
    "划转",
    "余额",
    "填报",
    "复核",
    "关联",
    "规范",
    "前置",
    "后置",
    "重要性",
    "安全",
    "说明",
    "参与角色",
    "输入项",
    "涉外",
    "报送",
    "表",
)

_ROLE_LABEL_DENY_SUBSTR = (
    "前置条件",
    "后置条件",
    "重要性",
    "数据安全",
    "其他说明",
    "参与角色",
    "无后置",
    "无前置",
)


def _name_touches_denylist(name: str) -> bool:
    n = (name or "").replace(" ", "").replace("\u3000", "")
    return any(w in n for w in _RESUME_NAME_DENY_SUBSTR)


def _looks_like_job_title_phrase(n: str) -> bool:
    """典型岗位/职称，不作自然人姓名（投标里常见「许海军-项目经理」右侧为职务）。"""
    t = (n or "").strip().replace("\u3000", "").replace(" ", "")
    if len(t) < 3:
        return False
    ends = (
        "工程师",
        "项目经理",
        "技术负责人",
        "项目总监",
        "技术总监",
        "产品经理",
        "系统架构师",
        "架构师",
        "设计师",
        "专家",
        "顾问",
        "负责人",
        "部经理",
        "部门经理",
    )
    if any(t.endswith(s) for s in ends):
        return True
    if len(t) <= 12 and t.endswith(("经理", "总监", "专员", "主管", "架构师")):
        return True
    return False


def _person_name_likely_requirement_section(s: str) -> bool:
    """章节/制度类长标题误进姓名字段（如「人员资质能力要求」「知识转移培训要求」）。"""
    t = re.sub(r"\s+", "", (s or "").strip())
    if len(t) > 10:
        return True
    if len(t) < 6:
        return False
    bad = (
        "管理",
        "培训",
        "要求",
        "资质",
        "罚则",
        "转移",
        "变更",
        "知识",
        "规模",
        "外包",
        "能力",
        "开发",
        "说明",
        "校验",
    )
    return any(w in t for w in bad)


def _is_plausible_cn_person_name(name: str) -> bool:
    """过滤需求/数据字典类标题被误抽成姓名。"""
    n = (name or "").strip().replace("\u3000", "").replace(" ", "")
    if len(n) < 2 or len(n) > 10:
        return False
    if n in _INVALID_EXACT_PERSON_NAME:
        return False
    if _looks_like_cn_id_card(n):
        return False
    if _name_touches_denylist(n):
        return False
    if _looks_like_job_title_phrase(n):
        return False
    if not re.fullmatch(r"[\u4e00-\u9fa5·]{2,10}", n):
        return False
    return True


# 切片 docx 内嵌图归属：按行识别「新人简历块起点」，须与 split_resume_into_persons / _person_header_match_items
# 支持的标题类型一致。若仅认「职务-姓名」连字符行，则「（一）项目总监：贾永超」无法推进 current_pi，
# current_pi 恒为 -1，学历/工作证明截图全部被丢弃。
_PERSON_LINE_ENUM_COLON_STRICT = re.compile(
    r"^\s*(?:[（(][一二三四五六七八九十百千]+[）)]\s*)?"
    r"(?P<role>[^\n\r：:]{2,48})"
    r"[：:]\s*"
    r"(?P<name>[\u4e00-\u9fa5·]{2,10})\s*$"
)

_PERSON_LINE_NUMBERED_NAME_ONLY = re.compile(
    r"^\s*\d+(?:\.\d+)+\.?\s+(?P<name>[\u4e00-\u9fa5·]{2,10})\s*$"
)

_PERSON_LINE_INLINE_HYPHEN = re.compile(
    r"^\s*(?P<role>\d{1,2}(?:\.\d{1,3})+\.?\s+[\u4e00-\u9fa5A-Za-z0-9（）()]{2,40})"
    r"[-－]"
    r"(?P<name>[\u4e00-\u9fa5·]{2,10})\s*$"
)


def _slice_img_person_boundary_line(s: str) -> bool:
    if not s or not s.strip():
        return False
    s = s.strip()
    if _PERSON_LINE_HEADER.match(s):
        return True
    m = _PERSON_LINE_ENUM_COLON_STRICT.match(s)
    if m:
        name_h = (m.group("name") or "").strip()
        return bool(name_h and _is_plausible_cn_person_name(name_h))
    m = _PERSON_LINE_INLINE_HYPHEN.match(s)
    if m:
        name_h = (m.group("name") or "").strip()
        return bool(name_h and _is_plausible_cn_person_name(name_h))
    m = _PERSON_LINE_NUMBERED_NAME_ONLY.match(s)
    if m:
        name_h = (m.group("name") or "").strip()
        return bool(name_h and _is_plausible_cn_person_name(name_h))
    return False


def _resolve_hyphen_role_and_name(raw_left: str, raw_right: str) -> tuple[str, str]:
    """
    连字符一行可能是「职务-姓名」或「姓名-职务」。
    返回 (role_label 已 sanitize, name_hint) 供 extract_person_structured 使用。
    """
    a0 = re.sub(r"\s+", "", (raw_left or "").strip())
    b0 = re.sub(r"\s+", "", (raw_right or "").strip())
    duty_a = _looks_like_job_title_phrase(a0)
    duty_b = _looks_like_job_title_phrase(b0)
    pa = _is_plausible_cn_person_name(a0)
    pb = _is_plausible_cn_person_name(b0)
    if duty_a and pb and not duty_b:
        return (_sanitize_header_role(a0) or raw_left.strip())[:120], b0[:10]
    if duty_b and pa and not duty_a:
        return (_sanitize_header_role(b0) or raw_right.strip())[:120], a0[:10]
    return (_sanitize_header_role(a0) or raw_left.strip())[:120], b0[:10]


def _postprocess_person_entry_row(row: dict) -> None:
    """入库前：职务/姓名对调修正；去掉明显章节标题误抽的姓名。"""
    pn = (row.get("person_name") or "").strip()
    rl = (row.get("role_label") or "").strip()
    pn_c = re.sub(r"\s+", "", pn)
    rl_c = re.sub(r"\s+", "", rl)
    if pn_c and rl_c:
        if _looks_like_job_title_phrase(pn_c) and _is_plausible_cn_person_name(rl_c) and not _looks_like_job_title_phrase(rl_c):
            row["person_name"] = rl[:80] if rl else None
            row["role_label"] = ((_sanitize_header_role(pn_c) or pn)[:120] if pn else None)
    pn = (row.get("person_name") or "").strip()
    if pn and _person_name_likely_requirement_section(pn):
        row["person_name"] = None
    elif pn and len(re.sub(r"[\s·•.\-]", "", pn)) > 10:
        row["person_name"] = None


def _role_label_too_technical(role: str) -> bool:
    r = (role or "").replace(" ", "").replace("\u3000", "")
    if len(r) > 36:
        return True
    return any(w in r for w in _ROLE_LABEL_DENY_SUBSTR)


def _chunk_has_resume_context(t: str) -> bool:
    """正文里出现典型「人员简历」语境，而非纯需求/数据表。"""
    t = t or ""
    markers = (
        "工作履历",
        "工作简历",
        "工作经验",
        "人员简历",
        "个人简历",
        "主要经历",
        "学历证书",
        "学历证明",
        "学位证书",
        "工作履历承诺",
        "项目经验证明",
        "简历证明材料",
    )
    return any(m in t for m in markers)


def _numbered_chunk_has_resume_table(chunk: str) -> bool:
    """编号小节标题后须跟典型简历表头/语境，避免 6.2.1 技术方案 等误当人。"""
    # Word 表格转文本多为「姓名\\t值」；须同时支持冒号类版式
    _kv = r"\s*[：:\t]"
    head = (chunk or "")[:1400]
    if re.search(r"姓名" + _kv, head):
        return True
    if re.search(r"性别" + _kv, head):
        return True
    if re.search(r"出生年月" + _kv, head):
        return True
    if re.search(r"学历" + _kv, head):
        return True
    if re.search(r"毕业学校" + _kv, head) or re.search(r"毕业院校" + _kv, head):
        return True
    return _chunk_has_resume_context(head)


def _should_keep_person_row(row: dict, chunk: str) -> bool:
    """
    去掉「项目团队人员情况」等大节里误命中「角色-姓名」行的需求/表结构碎片。
    保留：姓名像真人名，且（有联系方式/基本履历字段/或块内明显简历语境）。
    """
    name = (row.get("person_name") or "").strip()
    if not name or not _is_plausible_cn_person_name(name):
        return False
    rl = (row.get("role_label") or "").strip()
    if rl and _role_label_too_technical(rl):
        return False
    phone = (row.get("phone") or "").strip()
    email = (row.get("email") or "").strip()
    gender = (row.get("gender") or "").strip()
    age = (row.get("age") or "").strip()
    edu = (row.get("education_level") or "").strip()
    school = (row.get("graduate_school") or "").strip()
    wy = (row.get("work_years_hint") or "").strip()
    if phone or email:
        return True
    if gender and (age or edu):
        return True
    if edu and school:
        return True
    if wy and (gender or edu or age):
        return True
    if _chunk_has_resume_context(chunk):
        return True
    return False

_CT_EXT = {
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/jpg": "jpg",
    "image/gif": "gif",
    "image/bmp": "bmp",
    "image/tiff": "tiff",
    "image/x-emf": "emf",
    "image/x-wmf": "wmf",
}


def _classify_attachment_context(ctx: str) -> str:
    """根据附近正文关键词，把图片粗分为学历学位类 / 工作证明类 / 其它。"""
    c = (ctx or "").replace(" ", "").replace("\u3000", "")
    # 优先看「紧挨图片前」的一小段，避免全文里的「项目名称」干扰学历证截图
    # 简历块较长时，证明标题与图片之间可能隔多段文字，窗口略放大
    tail = c[-2800:] if len(c) > 2800 else c
    edu_tail = (
        "学历学位证书",
        "学历证书",
        "学历证明",
        "学位证书",
        "学历材料",
        "学历证明材料",
        "毕业证",
        "学位证",
        "受教育情况",
        "学历学位",
        "项目经理学历证书",
    )
    wk_tail = (
        "项目经验证明",
        "工作履历承诺",
        "合同证明材料",
        "工作证明",
        "业绩证明",
        "项目证明",
        "参与监管机构",
        "参与银行机构",
        "简历证明材料",
        "工作履历承诺函",
        "对应页码",
        "（一）参与监管机构",
        "（二）参与银行机构",
    )
    if any(k in tail for k in edu_tail):
        return "education"
    if any(k in tail for k in wk_tail):
        return "work"

    edu_kw = tuple(dict.fromkeys(edu_tail))  # 去重保持顺序
    wk_kw = (
        "项目经验证明",
        "工作履历承诺",
        "合同证明材料",
        "工作证明",
        "业绩证明",
        "项目证明",
        "参与监管机构",
        "参与银行机构",
        "简历证明材料",
        "工作履历承诺函",
    )
    ed = sum(1 for k in edu_kw if k in c)
    wk = sum(1 for k in wk_kw if k in c)
    if "工作经验" in tail or "工作简历" in tail:
        wk += 2
    if "项目名称" in tail or "项目角色" in tail or "责任描述" in tail:
        wk += 2
    if ed > wk:
        return "education"
    if wk > ed:
        return "work"
    if ed and wk and ed == wk:
        if "学历" in tail or "学位" in tail:
            return "education"
        return "work"
    return "other"


def _ext_for_image_part(part) -> str:
    ct = getattr(part, "content_type", "") or ""
    if ct in _CT_EXT:
        return _CT_EXT[ct]
    pn = str(getattr(part, "partname", "") or "")
    if "." in pn:
        return pn.rsplit(".", 1)[-1].lower()[:8]
    return "bin"


def _collect_image_rids_from_oxml(elem) -> list[str]:
    """段落或表格子树内，按 iter 深度优先顺序收集 a:blip / v:imagedata 的 rId。"""
    from docx.oxml.ns import qn

    rids: list[str] = []
    TAG_BLIP = qn("a:blip")
    for node in elem.iter():
        if node.tag == TAG_BLIP:
            rid = node.get(qn("r:embed")) or node.get(qn("r:link"))
            if rid and not str(rid).startswith("http"):
                rids.append(rid)
            continue
        if node.tag.endswith("imagedata"):
            rid = node.get(qn("r:id")) or node.get(qn("r:embed"))
            if rid and not str(rid).startswith("http"):
                rids.append(rid)
    return rids


def extract_resume_slice_images(
    split_path: str,
    person_entries: list[dict],
    work_dir: Path,
    record_index: int,
) -> None:
    """
    从二级切片 docx 中按文档顺序提取内嵌图，按「角色-姓名」标题归属到 person_entries，
    并写入 education_cert_image_paths / work_proof_image_paths / other_image_paths（绝对路径字符串列表）。
    """
    from docx import Document
    from lxml import etree

    get_elem_text = _splitter.get_elem_text

    for pe in person_entries:
        pe["education_cert_image_paths"] = []
        pe["work_proof_image_paths"] = []
        pe["other_image_paths"] = []

    if not person_entries:
        return

    n = len(person_entries)
    doc = Document(split_path)
    body = doc.element.body
    body_elems = [c for c in body if etree.QName(c.tag).localname in ("p", "tbl")]

    current_pi = -1
    recent_chunks: list[str] = []
    seq = 0
    attach_base = work_dir / "person_attachments" / f"rec_{record_index:02d}"
    attach_base.mkdir(parents=True, exist_ok=True)

    def _blob_for_rid(rid: str) -> tuple[bytes | None, str]:
        try:
            rel = doc.part.rels[rid]
        except KeyError:
            return None, "bin"
        part = rel.target_part
        blob = getattr(part, "blob", None)
        if not blob:
            return None, "bin"
        return blob, _ext_for_image_part(part)

    for elem in body_elems:
        txt = get_elem_text(elem).strip()
        if txt:
            for line in txt.split("\n"):
                s = line.strip()
                if s and _slice_img_person_boundary_line(s):
                    current_pi += 1
            if current_pi >= n:
                current_pi = n - 1
            recent_chunks.append(txt)
            while len(recent_chunks) > 20:
                recent_chunks.pop(0)

        ctx = "\n".join(recent_chunks)
        cat = _classify_attachment_context(ctx)

        for rid in _collect_image_rids_from_oxml(elem):
            blob, ext = _blob_for_rid(rid)
            if not blob:
                continue
            if current_pi < 0:
                continue
            seq += 1
            fn = f"p{current_pi:02d}_{cat}_{seq:04d}.{ext}"
            out = attach_base / fn
            out.write_bytes(blob)
            path_str = str(out.resolve())
            tgt = person_entries[current_pi]
            if cat == "education":
                tgt["education_cert_image_paths"].append(path_str)
            elif cat == "work":
                tgt["work_proof_image_paths"].append(path_str)
            else:
                tgt["other_image_paths"].append(path_str)


def _field(pattern: str, text: str, flags: int = 0) -> str | None:
    m = re.search(pattern, text, flags)
    return m.group(1).strip() if m else None


# 姓名与下一列/小节之间无制表符时，常见把「吴岫峰工作履历」「穆沅学历证书」写进同一单元格
_PERSON_NAME_STOP_KWS: tuple[str, ...] = (
    "学历证明材料",
    "学历学位证书",
    "学历证书",
    "学位证书",
    "学历证明",
    "证明材料",
    "工作履历承诺函",
    "工作履历",
    "工作简历",
    "人员简历",
    "承诺书",
    "毕业学校",
    "毕业院校",
    "性别",
    "年龄",
    "出生年月",
    "出生日期",
    "职称",
    "学历",
    "专业",
    "职务",
    "手机",
    "邮箱",
    "籍贯",
    "民族",
    "工作年限",
    "拟任职务",
    "身份证号",
    "身份证",
)
_FIELD_PERSON_NAME_RE = re.compile(
    r"姓名\s*[：:\t]\s*(?P<n>.+?)(?=(?:\t|\s|\u3000)*(?:"
    + "|".join(re.escape(k) for k in _PERSON_NAME_STOP_KWS)
    + r")(?:\s*[：:])?|\n|\Z)",
)


def _strip_person_name_doc_noise(s: str) -> str:
    """去掉误粘在姓名后的章节/材料标题（表格单元合并导出常见）。"""
    t = (s or "").strip()
    noise_suffixes = (
        "学历证明材料",
        "学历学位证书",
        "学历证书",
        "学位证书",
        "学历证明",
        "证明材料",
        "工作履历承诺函",
        "工作履历",
        "工作简历",
        "人员简历",
        "承诺书",
        "及业绩证明",
        "及业绩",
    )
    while t:
        hit = False
        for suf in noise_suffixes:
            if t.endswith(suf):
                t = t[: -len(suf)].strip()
                hit = True
                break
        if not hit:
            break
    return t


def _field_person_name(t: str) -> str | None:
    """
    Word 表格转纯文本时常为「姓名\\t张\\t辰\\t性别」——姓名单元被拆成多格，
    在下一列表头或「工作履历/学历证书」等小节词前截止；再剥常见误粘后缀。
    """
    m = _FIELD_PERSON_NAME_RE.search(t)
    if not m:
        return _extract_name_near_xingming(t)
    raw = m.group("n")
    s = re.sub(r"[\t \u3000]+", "", raw).strip()
    s = _strip_person_name_doc_noise(s)
    if not s:
        return _extract_name_near_xingming(t)
    if len(s) > 10:
        return _extract_name_near_xingming(t)
    if not _sanitize_person_name_string(s):
        return _extract_name_near_xingming(t)
    if re.fullmatch(r"[\u4e00-\u9fa5·•.\-A-Za-z]{1,40}", s):
        return s[:40]
    m2 = re.match(r"([\u4e00-\u9fa5·•.\-A-Za-z]{1,40})", s)
    cand = m2.group(1).strip()[:40] if m2 else None
    if cand and not _sanitize_person_name_string(cand):
        return _extract_name_near_xingming(t)
    return cand


def _sanitize_header_role(role: str) -> str:
    """
    Word 导出时表格单元常被拼成一行，「角色-姓名」前的 role 会夹带上一人简历里的噪声。
    优先截取常见职务关键词的最后一次出现；否则截短。
    """
    r = re.sub(r"\s+", "", (role or "").strip())
    if not r:
        return ""
    # 「3.2.1.项目经理」类：去掉章节编号，只保留职务用语
    r = re.sub(r"^\d+(?:\.\d+)*\.?", "", r)
    if not r:
        return ""
    noise = ("本科", "专科", "简历", "人员简历", "政治学与行政学")
    if len(r) <= 20 and not any(x in r for x in noise):
        return r[:120]
    # 从右向左找最后一次出现的职务类短语（投标简历里多写在 hyphen 前）
    role_keys = (
        "项目副经理",
        "项目经理",
        "技术负责人",
        "项目总监",
        "技术总监",
        "技术经理",
        "实施经理",
        "开发经理",
        "测试经理",
        "产品经理",
        "系统架构师",
        "架构师",
        "高级工程师",
        "工程师",
        "技术专家",
        "专家",
        "顾问",
        "负责人",
        "主管",
        "专员",
    )
    best_pos = -1
    best_kw = ""
    for kw in role_keys:
        p = r.rfind(kw)
        if p > best_pos:
            best_pos = p
            best_kw = kw
    if best_pos >= 0 and best_kw:
        return best_kw[:120]
    return r[:24]


def _source_docx_mtime_iso(path: str) -> str | None:
    """源 .docx 文件 mtime，UTC ISO，供库表区分同人多版文档。"""
    try:
        mt = os.path.getmtime(path)
        return (
            datetime.fromtimestamp(mt, tz=timezone.utc)
            .replace(microsecond=0)
            .isoformat()
            .replace("+00:00", "Z")
        )
    except OSError:
        return None


def _normalize_resume_text(t: str) -> str:
    """制表/全角符号粗归一，提高键值正则命中率；保留 \\t 列分隔，避免表格键值错位。"""
    if not t:
        return ""
    t = unicodedata.normalize("NFKC", t)
    t = t.replace("\r\n", "\n").replace("\r", "\n")
    t = t.replace("\u3000", " ").replace("\xa0", " ")
    return t


def _loose_scan_kv_in_text(z: str) -> dict[str, str]:
    """
    全文宽松扫描制表/混排简历中常见键值（与 _field 单行规则互补）。
    用于「学历证明」等切片正文不完整、键名略异（身份证号/证件号码）等场景。
    """
    if not z:
        return {}
    out: dict[str, str] = {}
    m = re.search(
        r"(?:身份证号码|身份证号|证件号码)\s*[：:\t]?\s*([0-9Xx\*]{15,22})",
        z,
        re.IGNORECASE,
    )
    if m:
        out["id_card"] = m.group(1).strip()[:24]
    m = re.search(
        r"(?:手机号|手机号码|移动电话|联系电话|办公电话|联系手机)\s*[：:\t]?\s*([0-9+\s\-–—]{6,26})",
        z,
    )
    if m:
        out["phone"] = re.sub(r"\s+", "", m.group(1).strip())[:50]
    m = re.search(
        r"(?:电子邮箱|E-mail|Email|电子信箱)\s*[：:\t]\s*([\w.\-+@]{4,120})",
        z,
        re.IGNORECASE,
    )
    if m:
        out["email"] = m.group(1).strip()[:120]
    m = re.search(
        r"(?:现所在单位|工作单位|所在单位|单位名称|任职单位|服务单位)\s*[：:\t]\s*([^\t\n]{2,200})",
        z,
    )
    if m:
        out["employer"] = m.group(1).strip().replace("\t", " ")[:200]
    m = re.search(r"学位\s*[：:\t]\s*([^\t\n\r]{1,24})", z)
    if m:
        out["degree"] = m.group(1).strip()[:40]
    m = re.search(r"籍贯\s*[：:\t]\s*([^\t\n\r]{2,40})", z)
    if m:
        out["native_place"] = m.group(1).strip()[:80]
    m = re.search(r"民族\s*[：:\t]\s*([^\t\n\r]{1,12})", z)
    if m:
        out["ethnicity"] = m.group(1).strip()[:40]
    m = re.search(r"职称\s*[：:\t]\s*([^\t\n\r]{1,40})", z)
    if m:
        out["professional_title"] = m.group(1).strip()[:80]
    m = re.search(
        r"拟在本项目担任职务\s*[：:\t]\s*([^\t\n\r]{2,80})", z
    )
    if m:
        out["proposed_project_role"] = m.group(1).strip()[:120]
    m = re.search(r"工作时间\s*[：:\t]\s*([^\t\n\r]{1,40})", z)
    if m:
        out["work_duration_text"] = m.group(1).strip()[:64]
    return out


def extract_person_structured(
    chunk: str, role_label: str = "", name_hint: str = ""
) -> dict:
    """从单人简历原文块抽取列字段（规则引擎，可后续换 AI）。"""
    t = _normalize_resume_text(chunk or "")
    name = _field_person_name(t)
    if not name:
        name = _field(r"姓名\s*[：:\t]\s*([^\s\n，,、;；]{1,24})", t)
        if name:
            name = re.sub(r"[\t ]+", "", name).strip()[:40] or None
    if not name:
        name = _field(r"人员姓名\s*[：:\t]\s*([^\n]+)", t)
        if name:
            name = re.sub(r"[\t\n ]+", "", name).split("，")[0].strip()[:40] or None
    nh_raw = (name_hint or "").strip()
    nh_compact = re.sub(r"\s+", "", nh_raw)
    if not name and nh_compact and _is_plausible_cn_person_name(nh_compact[:24]):
        name = nh_compact[:40]
    # 仅「表内姓名单格只取到一字」时用「角色-姓名」行补全（如 张 + 张辰），避免把整段标题并进来
    if (
        name
        and nh_compact
        and len(name) == 1
        and len(nh_compact) <= 12
        and nh_compact.startswith(name)
    ):
        name = nh_compact[:40]
    if name:
        name = _strip_person_name_doc_noise(name)[:40] or None
    name = _sanitize_person_name_string(name)
    if not name:
        name = _extract_name_near_xingming(t)

    gender = _field(r"性\s*别\s*[：:\t]\s*([^\t\n]{1,12})", t)
    birth_date = _field(
        r"(?:出生年月|出生日期)\s*[：:\t]\s*([^\t\n]{1,32})", t
    )
    age_only = _field(r"年龄\s*[：:\t]\s*([^\t\n]{1,12})", t)
    age = birth_date or age_only
    if not age:
        age = _field(r"(?:年龄|出生年月|出生日期)\s*[：:\t]\s*([^\t\n]{1,28})", t)
    edu = _field(
        r"(?:学历|文化程度|最高学历)\s*[：:\t]\s*([^\t\n]{1,40})", t
    )
    major = _field(r"专业\s*[：:\t]\s*([^\t\n]{1,80})", t)
    school = _field(
        r"(?:毕业学校|毕业院校|就读院校)\s*[：:\t]\s*([^\t\n]{1,120})",
        t,
    )
    wy = _field(
        r"(?:工作年限|本项目相关工作年限|相关行业工作年限)\s*[：:\t]\s*([^\t\n]{1,40})",
        t,
    )
    work_duration_text = _field(r"工作时间\s*[：:\t]\s*([^\t\n]{1,40})", t)
    if not wy and work_duration_text:
        wy = work_duration_text
    id_card = _field(
        r"(?:身份证号码|身份证号|证件号码)\s*[：:\t]\s*([0-9Xx\*]{15,22})", t
    )
    degree = _field(r"学位\s*[：:\t]\s*([^\t\n]{1,24})", t)
    employer = _field(
        r"(?:现所在单位|现所在机构或部门|工作单位|所在单位|单位名称|任职单位|服务单位)\s*[：:\t]\s*([^\t\n]{1,160})",
        t,
    )
    proposed_project_role = _field(
        r"(?:拟在本项目担任职务|拟在本项目担任中职务|拟任本项目职务|在本项目拟任职务)\s*[：:\t]\s*([^\t\n]{1,80})",
        t,
    )
    similar_project_exp = _field(
        r"同类项目工作经验\s*[：:\t]\s*([^\t\n]{1,500})", t
    )
    phone = _field(
        r"(?:手机|手机号|移动电话|联系电话|联系手机)\s*[：:\t]\s*([0-9+\s\-]{6,26})", t
    )
    email = _field(
        r"(?:邮箱|电子邮箱|E-mail|Email|电子信箱)\s*[：:\t]\s*([\w.\-+@]{4,120})",
        t,
        re.IGNORECASE,
    )
    # 制表排版里「职务」与「工作年限」常在同一行，勿用 [^\n] 以免把「工作年限」吃进职务
    duty = _field(
        r"(?:拟任职务|现任职务|职务|岗位|职位)\s*[：:\t]\s*([^\t\n]{1,40})", t
    )
    professional_title = _field(r"职称\s*[：:\t]\s*([^\t\n]{1,40})", t)
    ethnicity = _field(r"民族\s*[：:\t]\s*([^\t\n]{1,20})", t)
    native_place = _field(r"籍贯\s*[：:\t]\s*([^\t\n]{1,60})", t)

    project_experience = None
    # 优先「主要经历 / 工作履历」等小节，避免同块前文里出现「工作经验」误切片
    for kw in (
        "主要经历",
        "工作履历",
        "工作经验",
        "工作简历",
        "项目经验",
        "近两年业绩",
        "项目任职经历",
        "主要工作业绩",
        "同类项目经验",
    ):
        idx = t.find(kw)
        if idx == -1:
            continue
        tail = t[idx : idx + 15000]
        cut = tail.find("工作履历承诺函")
        if cut != -1:
            tail = tail[:cut]
        cut2 = tail.find("致 贵州银行")
        if cut2 != -1:
            tail = tail[:cut2]
        # 「主要经历 / 工作履历」后常跟「工作经验」项目列表，与 Word 小标题分段一致时只保留前者
        if kw in ("主要经历", "工作履历"):
            cut3 = tail.find("\n工作经验")
            if cut3 == -1:
                cut3 = tail.find("\n工作简历")
            if cut3 != -1:
                tail = tail[:cut3]
        project_experience = tail.strip() or None
        break

    role_clean = _sanitize_header_role(role_label) or None
    if role_clean:
        role_clean = role_clean[:120]
    extra = {}
    if duty:
        extra["duty_line"] = duty[:200]
    if professional_title:
        extra["professional_title"] = professional_title[:80]
    if ethnicity:
        extra["ethnicity"] = ethnicity[:40]
    if native_place:
        extra["native_place"] = native_place[:80]
    if birth_date:
        extra["birth_date"] = birth_date[:40]
    if age_only:
        extra["age_years"] = age_only[:12]
    if name_hint and not name:
        extra["name_from_header"] = re.sub(r"\s+", "", name_hint)[:40]
    if id_card:
        extra["id_card"] = id_card[:24]
    if degree:
        extra["degree"] = degree[:40]
    if employer:
        extra["employer"] = employer[:200]
    if proposed_project_role:
        extra["proposed_project_role"] = proposed_project_role[:120]
    if work_duration_text and work_duration_text != wy:
        extra["work_duration_text"] = work_duration_text[:64]
    if similar_project_exp:
        extra["similar_project_experience"] = similar_project_exp[:800]

    loose = _loose_scan_kv_in_text(t.replace("\r", "\n"))
    if loose.get("id_card") and not id_card:
        id_card = loose["id_card"]
    if loose.get("phone") and not phone:
        phone = loose["phone"]
    if loose.get("email") and not email:
        email = loose["email"]
    if loose.get("employer") and not employer:
        employer = loose["employer"]
    if loose.get("degree") and not degree:
        degree = loose["degree"]
    if loose.get("native_place") and not native_place:
        native_place = loose["native_place"]
    if loose.get("ethnicity") and not ethnicity:
        ethnicity = loose["ethnicity"]
    if loose.get("professional_title") and not professional_title:
        professional_title = loose["professional_title"]
    if loose.get("proposed_project_role") and not proposed_project_role:
        proposed_project_role = loose["proposed_project_role"]
    if loose.get("work_duration_text") and not work_duration_text:
        work_duration_text = loose["work_duration_text"]
        if not wy:
            wy = work_duration_text
    if id_card:
        extra["id_card"] = id_card[:24]
    if degree:
        extra["degree"] = degree[:40]
    if employer:
        extra["employer"] = employer[:200]
    if professional_title:
        extra["professional_title"] = professional_title[:80]
    if native_place:
        extra["native_place"] = native_place[:80]
    if ethnicity:
        extra["ethnicity"] = ethnicity[:40]
    if proposed_project_role:
        extra["proposed_project_role"] = proposed_project_role[:120]
    if work_duration_text and work_duration_text != (wy or ""):
        extra["work_duration_text"] = work_duration_text[:64]

    return {
        "role_label": role_clean,
        "person_name": (name or "")[:80] or None,
        "gender": (gender or "")[:20] or None,
        "age": (age or "")[:32] or None,
        "education_level": (edu or "")[:80] or None,
        "major": (major or "")[:200] or None,
        "graduate_school": (school or "")[:200] or None,
        "work_years_hint": (wy or "")[:64] or None,
        "id_card": (id_card or "")[:24] or None,
        "degree": (degree or "")[:40] or None,
        "employer": (employer or "")[:200] or None,
        "proposed_project_role": (proposed_project_role or "")[:120] or None,
        "work_duration_text": (work_duration_text or "")[:64] or None,
        "phone": (phone or "")[:50] or None,
        "email": (email or "")[:120] or None,
        "project_experience": project_experience,
        "raw_fragment": t[:50000] if t else None,
        "structured_json": extra if extra else None,
    }


def _person_header_match_items(block: str) -> list[tuple[re.Match, str]]:
    """
    合并「职务-姓名」「（一）总监：姓名」「编号+职务-姓名（行内）」「多级编号 + 姓名」「块首单独姓名行」等匹配。
    返回 [(match, kind), ...]，kind 含 hyphen | enum | inline | num | lead。
    """
    items: list[tuple[int, str, re.Match]] = []
    for m in _PERSON_HEADER.finditer(block):
        items.append((m.start(), "hyphen", m))
    for m in _PERSON_HEADER_ENUM_COLON.finditer(block):
        name_h = m.group("name").strip()
        if not name_h or not _is_plausible_cn_person_name(name_h):
            continue
        items.append((m.start(), "enum", m))
    for m in _PERSON_HEADER_INLINE.finditer(block):
        items.append((m.start(), "inline", m))
    for m in _PERSON_HEADER_NUMBERED.finditer(block):
        name_h = m.group("name").strip()
        if not name_h or not _is_plausible_cn_person_name(name_h):
            continue
        items.append((m.start(), "num", m))
    m_lead = _PERSON_HEADER_LEADING_NAME.match(block)
    if m_lead:
        nh = m_lead.group("name").strip()
        if nh and _is_plausible_cn_person_name(nh):
            items.append((m_lead.start(), "lead", m_lead))
    def _prio(k: str) -> int:
        return {"hyphen": 0, "enum": 0, "inline": 1, "num": 2, "lead": 3}.get(k, 9)

    by_start: dict[int, tuple[str, re.Match]] = {}
    for start, kind, m in sorted(items, key=lambda x: (x[0], _prio(x[1]))):
        if start in by_start:
            continue
        by_start[start] = (kind, m)
    ordered = sorted(by_start.items(), key=lambda kv: kv[0])
    return [(m, kind) for _start, (kind, m) in ordered]


def split_resume_into_persons(block: str) -> list[dict]:
    """
    按「职务-姓名」「编号+职务-姓名（含行内 \\t 分隔多段）」「编号 + 姓名」「块首姓名行」等切分为多人块，再对每块做字段抽取。
    若无此类标题则返回空列表（由上层保留章节级一条记录即可）。
    对明显「需求/数据表」误命中行做过滤，避免批量入库垃圾人名。
    """
    block = block or ""
    work = block.lstrip("\n\r \t\u3000")
    if not work.strip():
        return []
    match_items = _person_header_match_items(work)
    if not match_items:
        return []
    persons: list[dict] = []
    for i, (m, kind) in enumerate(match_items):
        start = m.start()
        end = match_items[i + 1][0].start() if i + 1 < len(match_items) else len(work)
        chunk = work[start:end].strip("\n")
        if kind in ("num", "lead"):
            if not _numbered_chunk_has_resume_table(chunk):
                continue
            role = ""
            name_h = m.group("name").strip()
        elif kind == "enum":
            raw_role = (m.group("role") or "").strip()
            role = re.sub(
                r"^[（(][一二三四五六七八九十百千]+[）)]\s*", "", raw_role
            ).strip()
            sr = _sanitize_header_role(role)
            role = ((sr or role) or "")[:120]
            name_h = (m.group("name") or "").strip()
        else:
            rs = (m.group("role") or "").strip()
            if (
                kind == "hyphen"
                and rs
                and re.match(r"^[一二三四五六七八九十百]+[、.．]", rs)
            ):
                if not any(
                    k in rs
                    for k in (
                        "经理",
                        "总监",
                        "工程师",
                        "专家",
                        "顾问",
                        "负责",
                        "主管",
                        "专员",
                        "架构师",
                        "设计师",
                    )
                ):
                    continue
            role, name_h = _resolve_hyphen_role_and_name(
                m.group("role").strip(), m.group("name").strip()
            )
            if role and _role_label_too_technical(role):
                continue
        row = extract_person_structured(chunk, role_label=role, name_hint=name_h)
        _postprocess_person_entry_row(row)
        if _should_keep_person_row(row, chunk):
            enrich_person_row(row, chunk)
            persons.append(row)
    return persons


def _safe_dir_segment(s: str, max_len: int = 200) -> str:
    """去掉 Windows 非法路径字符，便于作目录名。"""
    t = (s or "").strip()
    for ch in '<>:"/\\|?*':
        t = t.replace(ch, "_")
    t = re.sub(r"\s+", " ", t).strip()
    return t[:max_len] if t else ""


def resolve_output_folder_base_name(
    docx_path: str,
    *,
    folder_base_from: str = "parent",
    folder_base_override: str = "",
) -> str:
    """
    输出子目录「主名」段（不含 __标签_批次）。

    folder_base_override 非空时优先使用（经安全化截断）。
    folder_base_from:
      - parent: docx 父目录名（项目文件夹），无效时退回 stem
      - stem: 仅 docx 主文件名
      - both: 父目录名__主文件名（二者不同且均有效时），否则等价于 parent
    """
    o = _safe_dir_segment(folder_base_override, 200)
    if o:
        return o
    p = Path(docx_path).resolve()
    stem = _safe_dir_segment(p.stem, 120)
    parent = _safe_dir_segment(p.parent.name, 200)
    mode = (folder_base_from or "parent").strip().lower()
    if mode == "stem":
        return stem or "resume"
    if mode == "both":
        if (
            parent
            and parent not in (".", "..")
            and stem
            and parent != stem
        ):
            return _safe_dir_segment(f"{parent}__{stem}", 200)
        if parent and parent not in (".", "..") and len(parent) >= 2:
            return parent
        return stem or "resume"
    if parent and parent not in (".", "..") and len(parent) >= 2:
        return parent
    return stem or "resume"


def probe_resume_headings(input_docx: str) -> dict:
    """
    轻量探测：是否含简历相关大纲（与 run_pipeline 相同的 H1/H2 关键词规则），不做拆分与入库。
    供批量任务先筛「最大 docx」再决定是否全量抽取。
    """
    p = Path(str(input_docx).strip()).resolve()
    if not p.is_file():
        return {"ok": False, "error": "file_not_found", "path": str(input_docx), "has_resume": False}
    suf = p.suffix.lower()
    if suf != ".docx":
        return {"ok": False, "error": "not_docx", "path": str(p), "has_resume": False}
    from docx import Document
    from lxml import etree

    doc = Document(str(p))
    body = doc.element.body
    body_elems = [c for c in body if etree.QName(c.tag).localname in ("p", "tbl")]
    h1_list = _splitter.find_heading_positions(body_elems, doc, target_level=0)
    h2_list = _splitter.find_heading_positions(body_elems, doc, target_level=1)
    h3_list = _splitter.find_heading_positions(body_elems, doc, target_level=2)
    h4_list = _splitter.find_heading_positions(body_elems, doc, target_level=3)
    resume_h1 = [text for _idx, text in h1_list if h1_matches_resume(text)]
    resume_h2 = [text for _idx, text in h2_list if h2_matches_resume(text)]
    resume_h3 = [text for _idx, text in h3_list if h2_matches_resume(text)]
    resume_h4 = [text for _idx, text in h4_list if h2_matches_resume(text)]
    has_resume = bool(resume_h1 or resume_h2 or resume_h3 or resume_h4)
    return {
        "ok": True,
        "path": str(p),
        "has_resume": has_resume,
        "matched_h1_titles": resume_h1[:40],
        "matched_h2_titles": resume_h2[:60],
        "matched_h3_titles": resume_h3[:60],
        "matched_h4_titles": resume_h4[:60],
    }


def run_pipeline(
    input_docx: str,
    out_root: Path,
    *,
    folder_base_from: str = "parent",
    folder_base: str = "",
    folder_tag: str = "简历解析",
    batch_id: str | None = None,
) -> dict:
    input_docx = str(Path(input_docx).resolve())
    src_mtime_iso = _source_docx_mtime_iso(input_docx)
    base_name = resolve_output_folder_base_name(
        input_docx,
        folder_base_from=folder_base_from,
        folder_base_override=folder_base,
    )
    tag = _safe_dir_segment(folder_tag, 40) or "简历解析"
    bid = (batch_id or "").strip() or datetime.now().strftime("%Y%m%d_%H%M%S")
    work = out_root / f"{base_name}__{tag}_{bid}"
    work.mkdir(parents=True, exist_ok=True)
    manifest_base = {"output_dir": str(work)}
    boundary_skip_titles: set[str] = set()

    # 先全量列一级标题：用空 sections 触发 splitter 的「按全部 H1 拆分」逻辑太重；
    # 这里直接复用 split_docx 前先取 H1 列表：读 Document + find_heading_positions
    from docx import Document
    from lxml import etree

    doc = Document(input_docx)
    body = doc.element.body
    body_elems = [c for c in body if etree.QName(c.tag).localname in ("p", "tbl")]
    h1_list = _splitter.find_heading_positions(body_elems, doc, target_level=0)
    h2_list = _splitter.find_heading_positions(body_elems, doc, target_level=1)
    h3_list = _splitter.find_heading_positions(body_elems, doc, target_level=2)
    h4_list = _splitter.find_heading_positions(body_elems, doc, target_level=3)

    resume_h1_sections = [
        {"title": text, "start_keyword": text}
        for _idx, text in h1_list
        if h1_matches_resume(text)
    ]

    resume_h2_sections = [
        {"title": text, "start_keyword": text}
        for _idx, text in h2_list
        if h2_matches_resume(text)
    ]
    resume_h3_sections = [
        {"title": text, "start_keyword": text}
        for _idx, text in h3_list
        if h2_matches_resume(text)
    ]
    resume_h4_sections = [
        {"title": text, "start_keyword": text}
        for _idx, text in h4_list
        if h2_matches_resume(text)
    ]

    heading_list_for_expand: list = []
    # 若一级未命中，则用二级标题在全文中按 outlineLevel=1 切片；再退回三、四级（不少标书把「人员资质」设为 Heading3）
    if resume_h1_sections:
        split_mode = "h1"
        resume_sections = resume_h1_sections
        outline_level = 0
    elif resume_h2_sections:
        split_mode = "h2_fallback"
        resume_sections = resume_h2_sections
        outline_level = 1
        heading_list_for_expand = h2_list
    elif resume_h3_sections:
        split_mode = "h2_fallback"
        resume_sections = resume_h3_sections
        outline_level = 2
        heading_list_for_expand = h3_list
    elif resume_h4_sections:
        split_mode = "h2_fallback"
        resume_sections = resume_h4_sections
        outline_level = 3
        heading_list_for_expand = h4_list
    else:
        split_mode = None
        resume_sections = []
        outline_level = 0

    # split_docx 在仅 1 个章节时会把 end_idx 设为全文末尾；为每个命中段追加「下一 H2」作边界
    if resume_sections and split_mode == "h2_fallback" and heading_list_for_expand:
        resume_sections, boundary_skip_titles = _expand_h2_sections_with_next_heading(
            heading_list_for_expand, resume_sections
        )

    manifest = {
        **manifest_base,
        "batch_id": bid,
        "source_docx": input_docx,
        "source_docx_modified_at": src_mtime_iso,
        "path_options": {
            "out_root": str(out_root.resolve()),
            "folder_base_from": folder_base_from,
            "folder_base": folder_base or None,
            "folder_base_resolved": base_name,
            "folder_tag": tag,
        },
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "split_mode": split_mode if resume_sections else None,
        "matched_h1_keywords": list(RESUME_H1_KEYWORDS),
        "matched_h2_extra_keywords": list(RESUME_H2_EXTRA_KEYWORDS),
        "h1_resume_titles": [s["title"] for s in resume_h1_sections],
        "h2_resume_titles": [s["title"] for s in resume_h2_sections],
        "h3_resume_titles": [s["title"] for s in resume_h3_sections],
        "h4_resume_titles": [s["title"] for s in resume_h4_sections],
        "boundary_skip_titles": sorted(boundary_skip_titles),
        "records": [],
    }

    if not resume_sections:
        manifest["warning"] = (
            "一至四级大纲标题均未命中简历/人员相关关键词；请检查大纲级别样式，"
            "或扩充 RESUME_H1_KEYWORDS / RESUME_H2_EXTRA_KEYWORDS。"
        )
        _write_manifest(work, manifest)
        return manifest

    slice_dir = work / ("h1_sections" if split_mode == "h1" else "h2_sections")
    r1 = split_docx(input_docx, str(slice_dir), resume_sections, outline_level)
    manifest["split_primary_debug"] = r1.get("debug")

    if not r1.get("success"):
        manifest["error"] = r1.get("error", "一级拆分失败")
        _write_manifest(work, manifest)
        return manifest

    for sec in r1.get("sections", []):
        title_h1 = sec.get("title") or ""
        if title_h1 in boundary_skip_titles:
            continue
        docx_h1 = sec.get("docx_path")
        text_h1 = sec.get("text") or ""
        if not docx_h1 or not Path(docx_h1).is_file():
            manifest["records"].append(
                {
                    "h1_section_title": title_h1,
                    "l2_section_title": None,
                    "split_docx_path": None,
                    "raw_text": text_h1,
                    "extracted": extract_fields_from_text(text_h1, title_h1),
                    "error": sec.get("error") or "无一级 docx",
                }
            )
            continue

        h2_dir = work / "h2_from_h1" / "".join(
            c for c in title_h1 if c.isalnum() or c in (" ", "-", "_", "（", "）")
        ).strip()[:60]
        h2_dir.mkdir(parents=True, exist_ok=True)

        # 切片内再按 outlineLevel=1 拆细（子文档内常见为「姓名 / 小节」）
        r2 = split_docx(docx_h1, str(h2_dir), [], 1)

        if not r2.get("success") or not r2.get("sections"):
            ex = extract_fields_from_text(text_h1, title_h1)
            manifest["records"].append(
                {
                    "h1_section_title": title_h1,
                    "l2_section_title": None,
                    "split_docx_path": docx_h1,
                    "raw_text": text_h1,
                    "extracted": ex,
                    "note": "二级拆分无结果，整段作为一条记录",
                }
            )
            continue

        for sub in r2["sections"]:
            t2 = sub.get("title") or ""
            p2 = sub.get("docx_path")
            txt = sub.get("text") or ""
            manifest["records"].append(
                {
                    "h1_section_title": title_h1,
                    "l2_section_title": t2,
                    "split_docx_path": p2,
                    "raw_text": txt,
                    "extracted": extract_fields_from_text(txt, t2),
                    "error": sub.get("error"),
                }
            )

    manifest["output_dir"] = str(work)
    for rec in manifest.get("records") or []:
        rec["person_entries"] = split_resume_into_persons(rec.get("raw_text") or "")
        if not rec.get("person_entries") and rec.get("raw_text"):
            ex = rec.get("extracted")
            if isinstance(ex, dict):
                enrich_section_extracted(ex, rec.get("raw_text") or "")

    for ri, rec in enumerate(manifest.get("records") or []):
        persons = rec.get("person_entries") or []
        sp = rec.get("split_docx_path")
        if not persons or not sp:
            continue
        pth = Path(str(sp))
        if not pth.is_file():
            continue
        try:
            extract_resume_slice_images(str(pth), persons, work, ri)
        except Exception as e_img:
            rec["image_extract_error"] = str(e_img)[:500]

    try:
        manifest["ai_enrich_enabled"] = bool(_enrich.is_enrich_enabled())
    except Exception:
        manifest["ai_enrich_enabled"] = False

    _write_manifest(work, manifest)
    return manifest


def _write_manifest(work: Path, manifest: dict) -> None:
    manifest.setdefault("output_dir", str(work))
    (work / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def main():
    ap = argparse.ArgumentParser(
        description="投标文件简历章节拆分与抽取",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    ap.add_argument(
        "input_docx",
        nargs="?",
        default="",
        help="原始响应文件 .docx（使用 --probe-docx 时本参数可省略）",
    )
    ap.add_argument(
        "--probe-docx",
        dest="probe_docx",
        default="",
        metavar="DOCX",
        help="仅探测该 docx 是否含简历相关大纲，打印 JSON 到 stdout 后退出（不做拆分）",
    )
    ap.add_argument(
        "--out",
        dest="out",
        default="",
        help="输出根目录；未传时读环境变量 RESUME_EXTRACT_OUT，仍空则用 ai_backend/data/resume-extract-output",
    )
    ap.add_argument(
        "--folder-base-from",
        dest="folder_base_from",
        choices=("parent", "stem", "both"),
        default=None,
        metavar="MODE",
        help="子目录主名来源（未传本参数时才读 RESUME_EXTRACT_FOLDER_BASE_FROM，再默认 parent）",
    )
    ap.add_argument(
        "--folder-base",
        dest="folder_base",
        default="",
        help="强制指定子目录主名（非空则覆盖 --folder-base-from）；可用环境变量 RESUME_EXTRACT_FOLDER_BASE",
    )
    ap.add_argument(
        "--folder-tag",
        dest="folder_tag",
        default="",
        help="主名与批次之间的标签，默认「简历解析」；可用 RESUME_EXTRACT_FOLDER_TAG",
    )
    ap.add_argument(
        "--batch-id",
        dest="batch_id",
        default="",
        help="自定义批次号（默认当前时间）；可用 RESUME_EXTRACT_BATCH_ID",
    )
    ap.add_argument(
        "--no-ai",
        dest="no_ai",
        action="store_true",
        help="禁用大模型补缺（与 RESUME_EXTRACT_USE_AI=0 相同）",
    )
    args = ap.parse_args()

    probe_path = (args.probe_docx or "").strip()
    if probe_path:
        out = probe_resume_headings(probe_path)
        print(json.dumps(out, ensure_ascii=False))
        sys.exit(0 if out.get("ok") else 2)

    doc_in = (args.input_docx or "").strip()
    if not doc_in:
        ap.error("请指定 input_docx，或使用 --probe-docx")
    args.input_docx = doc_in

    default_root = Path(__file__).resolve().parents[2] / "data" / "resume-extract-output"
    out_raw = (args.out or os.environ.get("RESUME_EXTRACT_OUT", "") or "").strip()
    out_root = Path(out_raw) if out_raw else default_root
    out_root.mkdir(parents=True, exist_ok=True)

    if args.folder_base_from is not None:
        fb_from = args.folder_base_from
    else:
        efb = (os.environ.get("RESUME_EXTRACT_FOLDER_BASE_FROM") or "").strip().lower()
        fb_from = efb if efb in ("parent", "stem", "both") else "parent"
    folder_base = (args.folder_base or os.environ.get("RESUME_EXTRACT_FOLDER_BASE", "") or "").strip()
    folder_tag = (args.folder_tag or os.environ.get("RESUME_EXTRACT_FOLDER_TAG", "") or "").strip() or "简历解析"
    batch_id = (args.batch_id or os.environ.get("RESUME_EXTRACT_BATCH_ID", "") or "").strip() or None

    if getattr(args, "no_ai", False):
        os.environ["RESUME_EXTRACT_USE_AI"] = "0"

    manifest = run_pipeline(
        doc_in,
        out_root,
        folder_base_from=fb_from,
        folder_base=folder_base,
        folder_tag=folder_tag,
        batch_id=batch_id,
    )
    print(json.dumps({"ok": True, "output_dir": manifest.get("output_dir")}, ensure_ascii=False))
    if manifest.get("warning"):
        print(json.dumps({"warning": manifest["warning"]}, ensure_ascii=False), file=sys.stderr)
    if manifest.get("error"):
        print(json.dumps({"error": manifest["error"]}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
